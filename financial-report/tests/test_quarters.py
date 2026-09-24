import copy,io,json,pathlib,sys,subprocess,unittest
import pandas as pd
from openpyxl import load_workbook
ROOT=pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/'scripts'))
from backend.data_validator import validate_dataset,validate_period
from backend.vnstock_connector import normalize_frame
from backend.excel_builder import ExcelBuilder
from backend.metrics import decorate
from refresh_data import refresh_order

class QuarterlyReportsTest(unittest.TestCase):
 def setUp(self):self.data=json.loads((ROOT/'data/VIC.json').read_text())['quarterly']
 def test_vic_matches_reviewed_june_2026_report(self):
  # Vingroup reviewed consolidated H1 2026, approved 29 Aug 2026.
  # PDF pages 9-16 (printed pages 6-13), million VND.
  sections={s['id']:{r['id']:r['values'] for r in s['rows']} for s in self.data['sections']}
  bs=sections['balance_sheet'];inc=sections['income_statement'];cf=sections['cash_flow']
  self.assertEqual(bs['total_assets']['2026-Q2'],1309346482)
  self.assertEqual(bs['cash_and_cash_equivalents']['2026-Q2'],76396702)
  for rid,total in [('sales',221968415),('net_sales',221920410),('net_accounting_profit_loss_before_tax',34400091),('net_profit_loss_after_tax',20904971),('attributable_to_parent_company',15987971)]:
   self.assertEqual(inc[rid]['2026-Q1']+inc[rid]['2026-Q2'],total)
  for rid,total in [('net_cash_inflows_outflows_from_operating_activities',49029541),('net_cash_inflows_outflows_from_investing_activities',-147456737),('net_cash_inflows_outflows_from_financing_activities',102334522),('net_increase_in_cash_and_cash_equivalents',3907326)]:
   self.assertEqual(cf[rid]['2026-Q1']+cf[rid]['2026-Q2'],total)
  self.assertEqual(cf['cash_and_cash_equivalents_at_the_beginning_of_period']['2026-Q2'],cf['cash_and_cash_equivalents_at_the_end_of_period']['2026-Q1'])
 def test_quarter_axes_do_not_alias_annual_or_missing_periods(self):
  self.assertEqual(validate_period(self.data,['2026-Q2','2025-Q3']),['2025-Q3','2026-Q2'])
  for p in [2026,'2026','2026-Q3','2026-Q5','2024-Q2']:
   with self.assertRaises(ValueError):validate_period(self.data,[p])
  with self.assertRaises(ValueError):validate_period(self.data,['2025-Q4'],['ratios'])
  bad=copy.deepcopy(self.data);bad['sections'][0]['rows'][0]['values']['2026']=3
  with self.assertRaises(ValueError):validate_dataset(bad)
 def test_duplicate_source_periods_are_not_arbitrarily_selected(self):
  frame=pd.DataFrame([{'item':'P/E','item_id':'pe','2026-Q2':12,'2025-Q4':8,'2025-Q4_1':25}])
  rows=normalize_frame(frame,'ratios','KBS','quarter')['rows']
  self.assertEqual(rows[0]['values'],{'2026-Q2':12})
 def test_computed_metrics_use_report_values_and_preserve_missing_inputs(self):
  d=decorate(self.data);rows={r['id']:r for r in d['sections'][-1]['rows']}
  self.assertAlmostEqual(rows['derived_gross_margin']['values']['2026-Q2'],29496057/117568392*100)
  self.assertAlmostEqual(rows['derived_net_margin']['values']['2026-Q2'],15294192/117568392*100)
  self.assertIsNone(rows['derived_roa_period']['values']['2025-Q3'])
  # No false growth figure without the same quarter of the prior year.
  self.assertNotIn('derived_revenue_growth',rows)
  altered=copy.deepcopy(self.data)
  inc=next(s for s in altered['sections'] if s['id']=='income_statement')
  next(r for r in inc['rows'] if r['id']=='net_sales')['values']['2026-Q2']=0
  rows={r['id']:r for r in decorate(altered)['sections'][-1]['rows']}
  self.assertIsNone(rows['derived_gross_margin']['values']['2026-Q2'])
 def test_quarter_source_date_mismatch_is_not_exported(self):
  # Actual source claims Q2/2026 gross margin -18.63%; the reviewed statements
  # imply 25.09%. Retain the source privately but block this labelled period.
  d=decorate(self.data);source=next(s for s in d['sections'] if s['id']=='ratios')
  self.assertGreater(source['periodChecks']['2026-Q2']['mismatched'],1)
  self.assertNotIn('2026-Q2',source['periods'])
  self.assertTrue(all(r['values'].get('2026-Q2') is None for r in source['rows']))
  self.assertTrue(any(r['values'].get('2026-Q2') is not None for r in source['rawRows']))
  with self.assertRaises(ValueError):ExcelBuilder().build_bytes(self.data,['2026-Q2'],['ratios'])
 def test_current_vn100_and_fair_refresh_order(self):
  companies=json.loads((ROOT/'data/companies.json').read_text());symbols=[c['symbol'] for c in companies]
  self.assertEqual(len(symbols),100);self.assertEqual(len(set(symbols)),100);self.assertIn('VIC',symbols)
  order=refresh_order(companies,{'VIC':{'checkedAt':'2026-09-23T10:00:00Z'},'MBB':{'checkedAt':'2026-09-23T08:00:00Z','status':'retained'}})
  self.assertEqual(order[-1],'VIC');self.assertLess(order.index('MBB'),order.index('VIC'));self.assertEqual(set(order),set(symbols))
 def test_quarter_excel_matches_browser_python_and_every_value(self):
  d=decorate(self.data);periods=['2025-Q3','2026-Q2'];reports=['balance_sheet','income_statement','cash_flow','derived_ratios']
  js="const fs=require('fs'),vm=require('vm');for(const f of process.argv.slice(1,3))vm.runInThisContext(fs.readFileSync(f,'utf8'));const d=FinancialMetrics.decorate(JSON.parse(fs.readFileSync(process.argv[3],'utf8')).quarterly);fs.writeFileSync(process.argv[4],FinancialXlsx.workbook(d,JSON.parse(process.argv[5]),JSON.parse(process.argv[6])));"
  target=pathlib.Path('/tmp/VIC_quarter_browser_test.xlsx')
  subprocess.run(['node','-e',js,str(ROOT/'frontend/metrics.js'),str(ROOT/'frontend/xlsx.js'),str(ROOT/'data/VIC.json'),str(target),json.dumps(periods),json.dumps(reports)],check=True)
  browser=load_workbook(target);python=load_workbook(io.BytesIO(ExcelBuilder().build_bytes(self.data,periods,reports)))
  self.assertEqual(browser.sheetnames,['BCTC','Chi_so'])
  for wb in [browser,python]:
   for ws in wb:
    self.assertEqual(ws.freeze_panes,'B5');self.assertEqual(ws['B4'].value,'Q3/2025');self.assertEqual(ws['C4'].value,'Q2/2026');pos=5
    sections=[s for s in d['sections'] if s['id'] in reports and (s['id']=='derived_ratios')==(ws.title=='Chi_so')]
    for section in sections:
     pos+=1
     for row in section['rows']:
      for col,p in enumerate(periods,2):
       value=row['values'].get(p);cell=ws.cell(pos,col)
       if value is None:self.assertIsNone(cell.value)
       else:self.assertAlmostEqual(value,cell.value,places=6);self.assertEqual(cell.data_type,'n')
      pos+=1
    self.assertTrue(all(c.comment is None for row in ws for c in row))
if __name__=='__main__':unittest.main()

import tempfile
import copy,io,json,pathlib,sys,unittest,subprocess
import pandas as pd
from openpyxl import load_workbook
ROOT=pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from backend.data_validator import available_years,validate_period
from backend.vnstock_connector import normalize_frame,vci_records_frame,DEEP_PERIOD_LIMITS
from backend.excel_builder import ExcelBuilder

class ReportsTest(unittest.TestCase):
 def setUp(self):self.data=json.loads((ROOT/'data/MBB.json').read_text())
 def test_verified_mbb(self):
  rows={r['id']:r for s in self.data['sections'] if s['id']!='ratios' for r in s['rows']}
  self.assertEqual(rows['mbb_101']['values']['2023'],3966)
  # Cash balance from audited 2022 statement, not an inferred number.
  cash=[r for r in rows.values() if 'cuối' in r['label'].lower() and 'tiền' in r['label'].lower()]
  self.assertTrue(any(r['values']['2022']==68223912 for r in cash))
  for y in map(str,range(2020,2026)):
   assets=[r['values'][y] for r in rows.values() if 'TỔNG TÀI SẢN' in r['label'].upper()]
   funding=[r['values'][y] for r in rows.values() if 'TỔNG NỢ PHẢI TRẢ VÀ VỐN CHỦ SỞ HỮU' in r['label'].upper()]
   if assets and funding:self.assertEqual(assets[0],funding[0])
 def test_company_specific_years_and_gaps(self):
  self.assertEqual(available_years(self.data),list(range(2020,2026)))
  with self.assertRaises(ValueError):validate_period(self.data,[2019])
  self.assertEqual(validate_period(self.data,[2020,2025,2020]),[2020,2025])
  with self.assertRaises(ValueError):validate_period(self.data,[2020],['ratios'])
  with self.assertRaises(ValueError):validate_period(self.data,[2022],['nonexistent'])
  d=copy.deepcopy(self.data)
  for s in d['sections']:
   if s['id']=='cash_flow':
    for r in s['rows']:r['values']['2023']=None
  with self.assertRaises(ValueError):validate_period(d,[2023])
 def test_vci_deep_records_are_not_truncated_to_four_periods(self):
  records=[{'year':y,'sales':y*10} for y in range(2026,2018,-1)]
  frame=vci_records_frame(records,{'sales':'Doanh thu'},'year')
  self.assertGreater(DEEP_PERIOD_LIMITS['year'],4)
  self.assertEqual(list(frame.columns),['item','item_id','2026','2025','2024','2023','2022','2021','2020','2019'])
  self.assertEqual(frame.iloc[0]['2019'],20190)

 def test_dashboard_defaults_to_latest_and_previous_period_with_raw_vci_ids(self):
  fixture={'symbol':'FPT','periodType':'year','periods':[2023,2024,2025],'sections':[
   {'id':'balance_sheet','basis':'point_in_time','rows':[
    {'id':'bsa53','label':'TỔNG CỘNG TÀI SẢN','unit':'triệu đồng','values':{'2023':100000,'2024':110000,'2025':121000}},
    {'id':'bsa54','label':'NỢ PHẢI TRẢ','unit':'triệu đồng','values':{'2023':60000,'2024':65000,'2025':70000}},
    {'id':'bsa78','label':'Vốn chủ sở hữu','unit':'triệu đồng','values':{'2023':40000,'2024':45000,'2025':51000}}]},
   {'id':'income_statement','basis':'year','rows':[
    {'id':'isa3','label':'Doanh thu thuần','unit':'triệu đồng','values':{'2023':50000,'2024':55000,'2025':66000}},
    {'id':'isa20','label':'Lãi/(lỗ) thuần sau thuế','unit':'triệu đồng','values':{'2023':5000,'2024':6000,'2025':9000}}]},
   {'id':'cash_flow','basis':'year','rows':[
    {'id':'cfa18','label':'Lưu chuyển tiền tệ ròng từ các hoạt động sản xuất kinh doanh','unit':'triệu đồng','values':{'2023':7000,'2024':8000,'2025':10000}}]}]}
  js="const D=require(process.argv[1]),d=JSON.parse(process.argv[2]);const a=D.model(d,[2023,2024,2025]);const b=D.model(d,[2023,2024,2025],'2025','2023');const q=D.comparisonPeriods({periods:['2025-Q4','2026-Q1','2026-Q2']});console.log(JSON.stringify({latest:a.latest,compare:a.compare,assets:a.metrics.find(x=>x.key==='assets'),revenue:a.metrics.find(x=>x.key==='revenue'),profit:a.metrics.find(x=>x.key==='profit'),cash:a.metrics.find(x=>x.key==='cash'),capital:a.capital,custom:b.compare,customDelta:b.metrics.find(x=>x.key==='assets').delta,q}));"
  out=subprocess.run(['node','-e',js,str(ROOT/'frontend/dashboard.js'),json.dumps(fixture)],check=True,capture_output=True,text=True)
  data=json.loads(out.stdout)
  self.assertEqual((data['latest'],data['compare']),('2025','2024'))
  self.assertEqual(data['assets']['value'],121);self.assertAlmostEqual(data['assets']['delta'],10)
  self.assertEqual(data['revenue']['value'],66);self.assertEqual(data['profit']['value'],9);self.assertEqual(data['cash']['value'],10)
  self.assertAlmostEqual(data['capital']['equityPercent'],51/121*100)
  self.assertEqual(data['custom'],'2023');self.assertAlmostEqual(data['customDelta'],21)
  self.assertEqual((data['q']['latest'],data['q']['prior']),('2026-Q2','2026-Q1'))

 def test_no_quarter_mislabeled_as_year(self):
  bad=pd.DataFrame([['Tỷ lệ','r',12,20]],columns=['item','item_id','2018','2018'])
  with self.assertRaises(ValueError):normalize_frame(bad,'ratios','VCI')
 def test_balance_totals_for_all_bundled_industries(self):
  count=0
  for symbol in ('FPT','VCB','HPG','VNM','ACB','TCB','SSI','BVH'):
   data=json.loads((ROOT/'data'/f'{symbol}.json').read_text())
   section=next(s for s in data['sections'] if s['id']=='balance_sheet')
   rows={r['id']:r['values'] for r in section['rows']}
   funding=next(rows[k] for k in ('total_resource','total_resources','liabilities_and_shareholders_equity','liabilities_and_shareholders_equities') if k in rows)
   for y in map(str,data['years']):
    self.assertIsNotNone(rows['total_assets'][y]);self.assertIsNotNone(funding[y])
    self.assertAlmostEqual(rows['total_assets'][y],funding[y],places=2,msg=f'{symbol} {y}');count+=1
  self.assertEqual(count,32)
 def test_zero_null_duplicate_and_scaling(self):
  df=pd.DataFrame([{'item':'Tài sản','item_id':'x','2025':1_200_000,'2024':None},{'item':'Dự phòng','item_id':'x','2025':0,'2024':-2_000_000}])
  d=normalize_frame(df,'balance_sheet','VCI')['rows']
  self.assertEqual(d[0]['values']['2025'],1.2);self.assertIsNone(d[0]['values']['2024']);self.assertEqual(d[1]['values']['2025'],0);self.assertEqual(d[1]['values']['2024'],-2);self.assertNotEqual(d[0]['id'],d[1]['id'])
 def test_share_quantity_keeps_its_own_unit(self):
  frame=pd.DataFrame([{'item':'Cổ phiếu đang lưu hành (Số lượng)','item_id':'outstanding_shares_volume','2025':2075914794}])
  row=normalize_frame(frame,'balance_sheet','VCI')['rows'][0]
  self.assertEqual(row['unit'],'cổ phiếu');self.assertEqual(row['values']['2025'],2075914794)
  data=json.loads((ROOT/'data/SSI.json').read_text());rows={r['id']:r for r in data['sections'][0]['rows']}
  outstanding=rows['outstanding_shares_volume']['values']['2025'];treasury=rows['treasury_stocks_volume']['values']['2025']
  self.assertEqual(outstanding,2075914794);self.assertEqual(treasury,1991468)
 def test_python_excel_numeric_cells_and_plain_format(self):
  wb=load_workbook(io.BytesIO(ExcelBuilder().build_bytes(self.data,[2020,2025],['balance_sheet','income_statement','cash_flow','off_balance'])))
  ws=wb['BCTC'];self.assertEqual(ws.freeze_panes,'B5');self.assertEqual(ws.cell(4,2).value,2020);self.assertEqual(ws.cell(4,3).value,2025)
  self.assertTrue(all(c.comment is None for row in ws for c in row))
  actual=[r[1].value for r in ws if r[0].value and 'Lãi cơ bản trên cổ phiếu' in str(r[0].value)]
  self.assertEqual(actual,[2993]);self.assertEqual(wb.sheetnames,['BCTC'])
 def test_js_export_matches_every_selected_numeric_value(self):
  for symbol,years,reports in [('MBB',[2020,2025],['balance_sheet','income_statement','cash_flow','off_balance']),('FPT',[2022,2024,2025],['balance_sheet','income_statement','cash_flow','ratios']),('VCB',[2025],['ratios'])]:
   data=json.loads((ROOT/'data'/f'{symbol}.json').read_text())
   js="const fs=require('fs'),vm=require('vm');vm.runInThisContext(fs.readFileSync(process.argv[1],'utf8'));const d=JSON.parse(fs.readFileSync(process.argv[2],'utf8'));fs.writeFileSync(process.argv[3],FinancialXlsx.workbook(d,JSON.parse(process.argv[4]),JSON.parse(process.argv[5])));"
   target=pathlib.Path(tempfile.gettempdir())/f'{symbol}_browser_test.xlsx'
   subprocess.run(['node','-e',js,str(ROOT/'frontend/xlsx.js'),str(ROOT/'data'/f'{symbol}.json'),str(target),json.dumps(years),json.dumps(reports)],check=True)
   wb=load_workbook(target);counter=0
   for sheet in wb:
    self.assertEqual(sheet.freeze_panes,'B5');pos=5
    sections=[s for s in data['sections'] if s['id'] in reports and ((sheet.title=='Chi_so')==(s['id']=='ratios'))]
    for s in sections:
     pos+=1
     for r in s['rows']:
      for c,y in enumerate(years,2):
       expected=r['values'].get(str(y));actual=sheet.cell(pos,c).value
       if expected is None:self.assertIsNone(actual)
       else:self.assertAlmostEqual(actual,expected,places=6);self.assertEqual(sheet.cell(pos,c).data_type,'n');counter+=1
      pos+=1
   self.assertGreater(counter,10)
if __name__=='__main__':unittest.main()

